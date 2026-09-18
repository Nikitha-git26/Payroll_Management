"""
Builds Payroll_Operations_Tracker.xlsx from the CSVs in /data.

7 tabs, live formulas (INDEX/MATCH lookups, SUMIFS/COUNTIFS aggregation,
nested IF SLA logic), and conditional formatting for blocked tasks and
payroll variance. Run scripts/seed_data.py and scripts/engine.py first.

Usage:
    python scripts/build_excel_workbook.py
"""

import datetime
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV_DIR = os.path.join(ROOT, "data")
OUT_PATH = os.path.join(ROOT, "Payroll_Operations_Tracker.xlsx")

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F2E88")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=10)
INPUT_FONT = Font(name=FONT_NAME, size=10, color="0000FF")
LABEL_FONT = Font(name=FONT_NAME, size=10, bold=True)
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color="1F2E88")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RED_FILL = PatternFill("solid", fgColor="F8D7DA")
YELLOW_FILL = PatternFill("solid", fgColor="FFF3CD")
GREEN_FILL = PatternFill("solid", fgColor="D4EDDA")
ORANGE_FILL = PatternFill("solid", fgColor="FCE4D6")


def load_csvs():
    customers = pd.read_csv(os.path.join(CSV_DIR, "customers.csv"))
    customers["onboarding_start_date"] = pd.to_datetime(customers["onboarding_start_date"]).dt.date
    customers["target_closure_date"] = pd.to_datetime(customers["target_closure_date"]).dt.date

    tasks = pd.read_csv(os.path.join(CSV_DIR, "onboarding_tasks.csv"))
    tasks["due_date"] = pd.to_datetime(tasks["due_date"]).dt.date

    return {
        "customers": customers,
        "employees": pd.read_csv(
            os.path.join(CSV_DIR, "employees.csv"),
            dtype={"pan_number": str, "uan_number": str, "bank_account_no": str},
        ),
        "tasks": tasks,
        "payroll_register": pd.read_csv(os.path.join(CSV_DIR, "payroll_register.csv")),
        "reconciliation": pd.read_csv(os.path.join(CSV_DIR, "payroll_reconciliation.csv")),
        "exceptions": pd.read_csv(os.path.join(CSV_DIR, "exception_report.csv")),
    }


def write_header(ws, headers, row=1):
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def autosize(ws, widths):
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w


def style_body_cell(cell, number_format=None):
    cell.font = BODY_FONT
    cell.border = BORDER
    if number_format:
        cell.number_format = number_format


def build_01_customer_tracker(wb, customers):
    ws = wb.create_sheet("01_Customer_Tracker")
    headers = [
        "customer_id", "company_name", "industry", "employee_count",
        "onboarding_start_date", "target_closure_date", "current_stage", "status",
        "total_tasks", "completed_tasks", "completion_pct", "blocked_tasks",
        "days_remaining", "sla_status",
    ]
    write_header(ws, headers)

    ws["P1"] = "Reference Date:"
    ws["P1"].font = LABEL_FONT
    ws["Q1"] = datetime.date(2026, 9, 18)
    ws["Q1"].font = INPUT_FONT
    ws["Q1"].number_format = "yyyy-mm-dd"

    n = len(customers)
    for i, row in customers.iterrows():
        r = i + 2
        ws.cell(r, 1, row.customer_id)
        ws.cell(r, 2, row.company_name)
        ws.cell(r, 3, row.industry)
        ws.cell(r, 4, row.employee_count)
        ws.cell(r, 5, row.onboarding_start_date).number_format = "yyyy-mm-dd"
        ws.cell(r, 6, row.target_closure_date).number_format = "yyyy-mm-dd"
        ws.cell(r, 7, row.current_stage)
        ws.cell(r, 8, row.status)
        ws.cell(r, 9, f"=COUNTIF('04_Task_Pipeline'!$B:$B,A{r})")
        ws.cell(r, 10, f"=COUNTIFS('04_Task_Pipeline'!$B:$B,A{r},'04_Task_Pipeline'!$F:$F,\"Completed\")")
        ws.cell(r, 11, f"=IF(I{r}=0,0,J{r}/I{r})").number_format = "0.0%"
        ws.cell(r, 12, f"=COUNTIFS('04_Task_Pipeline'!$B:$B,A{r},'04_Task_Pipeline'!$F:$F,\"Blocked\")")
        ws.cell(r, 13, f"=F{r}-$Q$1").number_format = "0"
        ws.cell(
            r, 14,
            f'=IF(G{r}="Completed","Completed",IF(L{r}>0,"Blocked",'
            f'IF(M{r}<0,"SLA Breach",IF(AND(M{r}<=7,K{r}<0.75),"At Risk",'
            f'IF(AND(M{r}<=14,K{r}<0.5),"At Risk","On Track")))))'
        )
        for c in range(1, 15):
            style_body_cell(ws.cell(r, c))

    autosize(ws, [10, 22, 16, 10, 16, 16, 16, 12, 10, 12, 11, 11, 11, 12])

    last = n + 1
    ws.conditional_formatting.add(
        f"H2:H{last}", CellIsRule(operator="equal", formula=['"Blocked"'], fill=RED_FILL)
    )
    ws.conditional_formatting.add(
        f"N2:N{last}",
        FormulaRule(formula=[f'OR($N2="Blocked",$N2="SLA Breach")'], fill=RED_FILL),
    )
    ws.conditional_formatting.add(
        f"N2:N{last}", CellIsRule(operator="equal", formula=['"At Risk"'], fill=YELLOW_FILL)
    )
    ws.conditional_formatting.add(
        f"N2:N{last}", CellIsRule(operator="equal", formula=['"On Track"'], fill=GREEN_FILL)
    )
    return ws


def build_02_employee_master(wb, employees):
    ws = wb.create_sheet("02_Employee_Master")
    headers = [
        "employee_id", "customer_id", "name", "department", "basic_pay", "hra",
        "special_allowance", "overtime_pay", "bonus", "bank_account_no",
        "pan_number", "uan_number", "attendance_days", "lop_days", "customer_name",
    ]
    write_header(ws, headers)

    for i, row in employees.iterrows():
        r = i + 2
        ws.cell(r, 1, row.employee_id)
        ws.cell(r, 2, row.customer_id)
        ws.cell(r, 3, row["name"])
        ws.cell(r, 4, row.department)
        ws.cell(r, 5, row.basic_pay).number_format = "#,##0"
        ws.cell(r, 6, row.hra).number_format = "#,##0"
        ws.cell(r, 7, row.special_allowance).number_format = "#,##0"
        ws.cell(r, 8, row.overtime_pay).number_format = "#,##0"
        ws.cell(r, 9, row.bonus).number_format = "#,##0"
        ws.cell(r, 10, str(row.bank_account_no) if pd.notna(row.bank_account_no) else "")
        ws.cell(r, 11, str(row.pan_number) if pd.notna(row.pan_number) else "")
        ws.cell(r, 12, str(row.uan_number) if pd.notna(row.uan_number) else "")
        ws.cell(r, 13, row.attendance_days)
        ws.cell(r, 14, row.lop_days)
        ws.cell(r, 15, f"=IFERROR(INDEX('01_Customer_Tracker'!$B:$B,MATCH(B{r},'01_Customer_Tracker'!$A:$A,0)),\"\")")
        for c in range(1, 16):
            style_body_cell(ws.cell(r, c))

    autosize(ws, [10, 11, 18, 12, 10, 9, 12, 11, 8, 14, 11, 13, 9, 8, 20])

    last = len(employees) + 1
    ws.conditional_formatting.add(f"J2:J{last}", CellIsRule(operator="equal", formula=['""'], fill=RED_FILL))
    ws.conditional_formatting.add(f"K2:K{last}", CellIsRule(operator="equal", formula=['""'], fill=ORANGE_FILL))
    ws.conditional_formatting.add(f"M2:M{last}", CellIsRule(operator="equal", formula=["0"], fill=YELLOW_FILL))
    ws.conditional_formatting.add(f"G2:G{last}", CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
    return ws


def build_03_payroll_register(wb, payroll):
    ws = wb.create_sheet("03_Payroll_Register")
    headers = [
        "employee_id", "customer_id", "gross_earnings", "expected_net_pay",
        "processed_net_pay", "pf_deduction", "esi_deduction", "pt_deduction",
        "tds_deduction", "variance", "department", "bank_account_no",
    ]
    write_header(ws, headers)

    for i, row in payroll.iterrows():
        r = i + 2
        ws.cell(r, 1, row.employee_id)
        ws.cell(r, 2, row.customer_id)
        ws.cell(r, 3, row.gross_earnings).number_format = "#,##0"
        ws.cell(r, 4, row.expected_net_pay).number_format = "#,##0"
        ws.cell(r, 5, row.processed_net_pay).number_format = "#,##0"
        ws.cell(r, 6, row.pf_deduction).number_format = "#,##0"
        ws.cell(r, 7, row.esi_deduction).number_format = "#,##0"
        ws.cell(r, 8, row.pt_deduction).number_format = "#,##0"
        ws.cell(r, 9, row.tds_deduction).number_format = "#,##0"
        ws.cell(r, 10, row.variance).number_format = "#,##0;(#,##0)"
        ws.cell(r, 11, f"=IFERROR(INDEX('02_Employee_Master'!$D:$D,MATCH(A{r},'02_Employee_Master'!$A:$A,0)),\"\")")
        ws.cell(r, 12, f"=IFERROR(INDEX('02_Employee_Master'!$J:$J,MATCH(A{r},'02_Employee_Master'!$A:$A,0)),\"\")")
        for c in range(1, 13):
            style_body_cell(ws.cell(r, c))

    autosize(ws, [10, 11, 13, 14, 15, 11, 11, 10, 11, 11, 14, 14])
    last = len(payroll) + 1
    ws.conditional_formatting.add(f"J2:J{last}", CellIsRule(operator="notEqual", formula=["0"], fill=RED_FILL))
    return ws


def build_04_task_pipeline(wb, tasks):
    ws = wb.create_sheet("04_Task_Pipeline")
    headers = ["task_id", "customer_id", "task_name", "owner", "due_date", "status", "company_name"]
    write_header(ws, headers)

    for i, row in tasks.iterrows():
        r = i + 2
        ws.cell(r, 1, row.task_id)
        ws.cell(r, 2, row.customer_id)
        ws.cell(r, 3, row.task_name)
        ws.cell(r, 4, row.owner)
        ws.cell(r, 5, row.due_date).number_format = "yyyy-mm-dd"
        ws.cell(r, 6, row.status)
        ws.cell(r, 7, f"=IFERROR(INDEX('01_Customer_Tracker'!$B:$B,MATCH(B{r},'01_Customer_Tracker'!$A:$A,0)),\"\")")
        for c in range(1, 8):
            style_body_cell(ws.cell(r, c))

    autosize(ws, [10, 12, 26, 13, 13, 13, 22])
    last = len(tasks) + 1
    ws.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="equal", formula=['"Blocked"'], fill=RED_FILL))
    ws.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="equal", formula=['"Pending"'], fill=YELLOW_FILL))
    ws.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="equal", formula=['"Completed"'], fill=GREEN_FILL))
    return ws


def build_05_exception_log(wb, exceptions):
    ws = wb.create_sheet("05_Exception_Log")
    headers = ["customer_id", "employee_id", "exception_type", "severity", "recommended_action", "company_name"]
    write_header(ws, headers)

    for i, row in exceptions.iterrows():
        r = i + 2
        ws.cell(r, 1, row.customer_id)
        ws.cell(r, 2, row.employee_id)
        ws.cell(r, 3, row.exception_type)
        ws.cell(r, 4, row.severity)
        ws.cell(r, 5, row.recommended_action)
        ws.cell(r, 6, f"=IFERROR(INDEX('01_Customer_Tracker'!$B:$B,MATCH(A{r},'01_Customer_Tracker'!$A:$A,0)),\"\")")
        for c in range(1, 7):
            style_body_cell(ws.cell(r, c))

    autosize(ws, [12, 12, 28, 10, 46, 22])
    last = len(exceptions) + 1
    ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(operator="equal", formula=['"Critical"'], fill=RED_FILL))
    ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(operator="equal", formula=['"High"'], fill=ORANGE_FILL))
    ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(operator="equal", formula=['"Medium"'], fill=YELLOW_FILL))
    return ws


def build_06_reconciliation_summary(wb, reconciliation):
    ws = wb.create_sheet("06_Reconciliation_Summary")
    headers = [
        "customer_id", "company_name", "expected_gross", "expected_net", "processed_net",
        "variance", "reconciliation_status", "sumifs_expected_net_check", "sumifs_processed_net_check",
    ]
    write_header(ws, headers)

    for i, row in reconciliation.iterrows():
        r = i + 2
        ws.cell(r, 1, row.customer_id)
        ws.cell(r, 2, f"=IFERROR(INDEX('01_Customer_Tracker'!$B:$B,MATCH(A{r},'01_Customer_Tracker'!$A:$A,0)),\"\")")
        ws.cell(r, 3, row.expected_gross).number_format = "#,##0"
        ws.cell(r, 4, row.expected_net).number_format = "#,##0"
        ws.cell(r, 5, row.processed_net).number_format = "#,##0"
        ws.cell(r, 6, row.variance).number_format = "#,##0;(#,##0)"
        ws.cell(r, 7, row.reconciliation_status)
        ws.cell(r, 8, f"=SUMIFS('03_Payroll_Register'!$D:$D,'03_Payroll_Register'!$B:$B,A{r})").number_format = "#,##0"
        ws.cell(r, 9, f"=SUMIFS('03_Payroll_Register'!$E:$E,'03_Payroll_Register'!$B:$B,A{r})").number_format = "#,##0"
        for c in range(1, 10):
            style_body_cell(ws.cell(r, c))

    autosize(ws, [11, 22, 14, 13, 14, 12, 18, 20, 20])
    last = len(reconciliation) + 1
    ws.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="notEqual", formula=["0"], fill=RED_FILL))
    ws.conditional_formatting.add(
        f"G2:G{last}", CellIsRule(operator="equal", formula=['"Unreconciled"'], fill=RED_FILL)
    )
    ws.conditional_formatting.add(
        f"G2:G{last}", CellIsRule(operator="equal", formula=['"Under Review"'], fill=YELLOW_FILL)
    )
    ws.conditional_formatting.add(
        f"G2:G{last}", CellIsRule(operator="equal", formula=['"Reconciled"'], fill=GREEN_FILL)
    )
    return ws


def build_07_kpi_dashboard(wb, n_customers):
    ws = wb.create_sheet("07_KPI_Dashboard")
    ws["B2"] = "PayFlow — Payroll Operations KPI Dashboard"
    ws["B2"].font = TITLE_FONT
    ws.merge_cells("B2:E2")

    last_c = n_customers + 1

    rows = [
        ("Onboarding Pipeline", None),
        ("Total Onboarding Accounts", "=COUNTA('01_Customer_Tracker'!$A$2:$A$%d)" % last_c),
        ("On Track", "=COUNTIF('01_Customer_Tracker'!$H:$H,\"On Track\")"),
        ("At Risk", "=COUNTIF('01_Customer_Tracker'!$H:$H,\"At Risk\")"),
        ("Blocked", "=COUNTIF('01_Customer_Tracker'!$H:$H,\"Blocked\")"),
        ("Completed", "=COUNTIF('01_Customer_Tracker'!$H:$H,\"Completed\")"),
        ("SLA Breach (recomputed)", "=COUNTIF('01_Customer_Tracker'!$N:$N,\"SLA Breach\")"),
        ("Average Task Completion %", "=AVERAGE('01_Customer_Tracker'!$K:$K)"),
        (None, None),
        ("Payroll Reconciliation", None),
        ("Reconciled Accounts", "=COUNTIF('06_Reconciliation_Summary'!$G:$G,\"Reconciled\")"),
        ("Under Review", "=COUNTIF('06_Reconciliation_Summary'!$G:$G,\"Under Review\")"),
        ("Unreconciled", "=COUNTIF('06_Reconciliation_Summary'!$G:$G,\"Unreconciled\")"),
        ("Total Expected Gross (₹)", "=SUM('06_Reconciliation_Summary'!$C:$C)"),
        ("Total Expected Net (₹)", "=SUM('06_Reconciliation_Summary'!$D:$D)"),
        ("Total Processed Net (₹)", "=SUM('06_Reconciliation_Summary'!$E:$E)"),
        ("Total Net Variance (₹)", "=SUM('06_Reconciliation_Summary'!$F:$F)"),
        (None, None),
        ("Exception Severity", None),
        ("Critical Exceptions", "=COUNTIF('05_Exception_Log'!$D:$D,\"Critical\")"),
        ("High Exceptions", "=COUNTIF('05_Exception_Log'!$D:$D,\"High\")"),
        ("Medium Exceptions", "=COUNTIF('05_Exception_Log'!$D:$D,\"Medium\")"),
        ("Total Exceptions Logged", "=COUNTA('05_Exception_Log'!$A$2:$A$100000)"),
    ]

    r = 4
    for label, formula in rows:
        if label is None:
            r += 1
            continue
        if formula is None:
            cell = ws.cell(r, 2, label)
            cell.font = Font(name=FONT_NAME, size=11, bold=True, color="1F2E88")
            r += 1
            continue
        ws.cell(r, 2, label).font = LABEL_FONT
        vcell = ws.cell(r, 3, formula)
        vcell.font = BODY_FONT
        if "%" in label:
            vcell.number_format = "0.0%"
        elif "₹" in label:
            vcell.number_format = "#,##0"
        for c in (2, 3):
            ws.cell(r, c).border = BORDER
        r += 1

    autosize(ws, [4, 30, 16])
    ws.sheet_view.showGridLines = False
    return ws


def main():
    data = load_csvs()

    wb = Workbook()
    wb.remove(wb.active)

    build_01_customer_tracker(wb, data["customers"])
    build_02_employee_master(wb, data["employees"])
    build_03_payroll_register(wb, data["payroll_register"])
    build_04_task_pipeline(wb, data["tasks"])
    build_05_exception_log(wb, data["exceptions"])
    build_06_reconciliation_summary(wb, data["reconciliation"])
    build_07_kpi_dashboard(wb, len(data["customers"]))

    wb.save(OUT_PATH)
    print(f"Workbook written to: {OUT_PATH}")


if __name__ == "__main__":
    main()
